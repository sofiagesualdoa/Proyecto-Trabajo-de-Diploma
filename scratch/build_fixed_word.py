# -*- coding: utf-8 -*-
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

INDEX_ENTRIES = [
    (0, "G01. Propósito y Motivación", "g01"),
    (0, "G02. Requisitos Funcionales", "g02"),
    (0, "G03. Definiciones, Acrónimos y Abreviaciones", "g03"),
    (0, "G04. Tabla de Roles", "g04"),
    (0, "G05. Otros Requisitos", "g05"),
    (0, "G06. Diagramas de Clases", "g06"),
    (0, "G07. Diagrama Entidad Relación", "g07"),
    (0, "N01. Especificación funcional por proceso de negocio", "n01"),
    (1, "N01.PN1.Gestión de ventas en tienda física", "n01pn1"),
    (1, "N01.PN2.Gestión de préstamos y devolución de libros en tienda física", "n01pn2"),
    (0, "N02. Especificaciones de Casos de Uso por Proceso de Negocio", "n02"),
    (1, "N02.PN1.1 Diagrama de casos de uso general", "n02pn11"),
    (1, "N02.PN1.2 CUN01 “Registrar libro en el carrito”", "n02pn12"),
    (1, "N02.PN1.3 CUN02 “Seleccionar libro”", "n02pn13"),
    (1, "N02.PN1.4 CUN03 “Cobrar venta”", "n02pn14"),
    (1, "N02.PN1.5 CUN04 “Registrar cliente”", "n02pn15"),
    (1, "N02.PN2.1 Diagrama de casos de uso general", "n02pn21"),
    (1, "N02.PN2.2 CUN05 “Registrar préstamo”", "n02pn22"),
    (1, "N02.PN2.3 CUN06 “Cobrar préstamo”", "n02pn23"),
    (1, "N02.PN2.4 CUN07 “Seleccionar préstamo vigente”", "n02pn26"),
    (1, "N02.PN2.5 CUN08 “Cobrar recargo”", "n02pn25"),
    (1, "N02.PN2.6 CUN09 “Registrar devolución”", "n02pn24"),
    (0, "T01. Arquitectura", "t01"),
    (1, "T01.1. Diseño de la Arquitectura", "t011"),
    (1, "T01.2. Configuración de la Arquitectura en Visual Studio", "t012"),
    (1, "T01.3. Diagrama de Secuencia de Persistencia de Datos", "t013"),
    (1, "T01.4. Diagrama de Secuencia de Consulta de Datos", "t014"),
    (1, "T01.5. Esquema Jerárquico", "t015"),
    (1, "T01.6. GUI (Menú Principal)", "t016"),
    (0, "T02. Usuarios", "t02"),
    (1, "T02.1. Diseño Crear Usuario", "t021"),
    (1, "T02.2. Diseño Login", "t022"),
    (2, "CU20: Administrar Error de Integridad", "cu20"),
    (1, "T02.3. Diseño Desbloquear Usuario", "t023"),
    (1, "T02.4. Diseño Modificar Usuario", "t024"),
    (1, "T02.5. Diseño Cambiar Clave", "t025"),
    (1, "T02.6. Diseño Logout", "t026"),
    (1, "T02.7. Diseño Activar/Desactivar Usuario", "t027"),
    (0, "T03. Encriptación", "t03"),
    (0, "T04. Perfiles", "t04"),
    (1, "T04.1. Crear Familia", "t041"),
    (1, "T04.2. Eliminar Familia", "t042"),
    (1, "T04.3. Crear Perfil", "t043"),
    (1, "T04.4. Eliminar Perfil", "t044"),
    (1, "T04.5. Agregar Permisos a Familia", "t045"),
    (1, "T04.6. Quitar Permisos a Familia", "t046"),
    (1, "T04.7. Agregar Permisos a Perfil", "t047"),
    (1, "T04.8. Quitar Permisos a Perfil", "t048"),
    (1, "T04.9. Agregar Familias a Perfil", "t049"),
    (1, "T04.10. Quitar Familias a Perfil", "t0410"),
    (0, "T05. Idiomas", "t05"),
    (0, "T06. Bitácora de Eventos", "t06"),
    (0, "T07. Gestión de Respaldos", "t07"),
    (0, "T08. Dígito Verificador", "t08"),
]

def create_index_paragraph(text, anchor, level=0):
    p = OxmlElement('w:p')
    pPr = OxmlElement('w:pPr')
    if level == 1:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), '720')
        pPr.append(ind)
    elif level == 2:
        ind = OxmlElement('w:ind')
        ind.set(qn('w:firstLine'), '1440')
        pPr.append(ind)
    
    p_rPr = OxmlElement('w:rPr')
    rStyle = OxmlElement('w:rStyle')
    rStyle.set(qn('w:val'), 'Hipervnculo')
    p_rPr.append(rStyle)
    pPr.append(p_rPr)
    p.append(pPr)
    
    hl = OxmlElement('w:hyperlink')
    hl.set(qn('w:anchor'), anchor)
    hl.set(qn('w:history'), '1')
    
    r = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    rStyle_r = OxmlElement('w:rStyle')
    rStyle_r.set(qn('w:val'), 'Hipervnculo')
    rPr.append(rStyle_r)
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '1155CC')
    rPr.append(color)
    u = OxmlElement('w:u')
    u.set(qn('w:val'), 'single')
    rPr.append(u)
    r.append(rPr)
    
    t = OxmlElement('w:t')
    t.text = text
    r.append(t)
    
    hl.append(r)
    p.append(hl)
    return p

def create_page_break_paragraph():
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    return p

def main():
    docx_path = r'c:\Users\usuario\OneDrive\Facultad\3° Año\2° Cuatrimestre\Trabajo de Diploma\Carpeta\Proyecto-Trabajo-de-Diploma\scratch\Proyecto.docx'
    out_path = r'c:\Users\usuario\OneDrive\Facultad\3° Año\2° Cuatrimestre\Trabajo de Diploma\Carpeta\Proyecto-Trabajo-de-Diploma\scratch\Proyecto_fixed.docx'

    print("Loading document...")
    doc = docx.Document(docx_path)
    body = doc._body._body

    # Locate p24 (Índice:), p25 (empty), and p67 (G01)
    p24 = doc.paragraphs[24]._p
    p25 = doc.paragraphs[25]._p
    p67 = doc.paragraphs[67]._p

    idx24 = body.index(p24)
    idx25 = body.index(p25)
    idx67 = body.index(p67)

    print(f"p24 index in body: {idx24}, text: {doc.paragraphs[24].text}")
    print(f"p25 index in body: {idx25}, text: {doc.paragraphs[25].text}")
    print(f"p67 index in body: {idx67}, text: {doc.paragraphs[67].text}")

    # Verify p67 is indeed G01
    assert "G01. Propósito y Motivación" in doc.paragraphs[67].text, "p67 is not G01!"

    # Remove all elements between p25 and p67
    elements_to_remove = [body[i] for i in range(idx25 + 1, idx67)]
    print(f"Removing {len(elements_to_remove)} corrupted index elements...")
    for el in elements_to_remove:
        body.remove(el)

    # Insert new index entries right after p25
    insert_point = p25
    print(f"Inserting {len(INDEX_ENTRIES)} index entries...")
    for level, text, anchor in INDEX_ENTRIES:
        new_p = create_index_paragraph(text, anchor, level)
        insert_point.addnext(new_p)
        insert_point = new_p

    # Insert page break paragraph right before G01
    pb_p = create_page_break_paragraph()
    insert_point.addnext(pb_p)

    print("Saving modified document...")
    doc.save(out_path)
    print(f"Saved successfully to {out_path}!")

if __name__ == '__main__':
    main()
