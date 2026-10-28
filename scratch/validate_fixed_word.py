# -*- coding: utf-8 -*-
import docx

INDEX_ENTRIES = [
    (0, "G01. Propósito y Motivación", "g01"),
    (0, "G02. Requisitos Funcionales", "g02"),
    (0, "G03. Definiciones, Acrónimos y Abreviaciones", "g03"),
    (0, "G04. Tabla de Roles", "g04"),
    (0, "G05. Otros Requisitos", "g05"),
    (0, "G06. Diagramas de Clases", "g06"),
    (0, "G07. Diagrama Entidad Relación", "g07"),
    (0, "N01. Especificación funcional por proceso de negocio", "n01"),
    (1, "N01.PN1.Gestión de ventas en tienda física", "n01pn1"),
    (1, "N01.PN2.Gestión de préstamos y devolución de libros en tienda física", "n01pn2"),
    (0, "N02. Especificaciones de Casos de Uso por Proceso de Negocio", "n02"),
    (1, "N02.PN1.1 Diagrama de casos de uso general", "n02pn11"),
    (1, "N02.PN1.2 CUN01 “Registrar libro en el carrito”", "n02pn12"),
    (1, "N02.PN1.3 CUN02 “Seleccionar libro”", "n02pn13"),
    (1, "N02.PN1.4 CUN03 “Cobrar venta”", "n02pn14"),
    (1, "N02.PN1.5 CUN04 “Registrar cliente”", "n02pn15"),
    (1, "N02.PN2.1 Diagrama de casos de uso general", "n02pn21"),
    (1, "N02.PN2.2 CUN05 “Registrar préstamo”", "n02pn22"),
    (1, "N02.PN2.3 CUN06 “Cobrar préstamo”", "n02pn23"),
    (1, "N02.PN2.4 CUN07 “Seleccionar préstamo vigente”", "n02pn26"),
    (1, "N02.PN2.5 CUN08 “Cobrar recargo”", "n02pn25"),
    (1, "N02.PN2.6 CUN09 “Registrar devolución”", "n02pn24"),
    (0, "T01. Arquitectura", "t01"),
    (1, "T01.1. Diseño de la Arquitectura", "t011"),
    (1, "T01.2. Configuración de la Arquitectura en Visual Studio", "t012"),
    (1, "T01.3. Diagrama de Secuencia de Persistencia de Datos", "t013"),
    (1, "T01.4. Diagrama de Secuencia de Consulta de Datos", "t014"),
    (1, "T01.5. Esquema Jerárquico", "t015"),
    (1, "T01.6. GUI (Menú Principal)", "t016"),
    (0, "T02. Usuarios", "t02"),
    (1, "T02.1. Diseño Crear Usuario", "t021"),
    (1, "T02.2. Diseño Login", "t022"),
    (2, "CU20: Administrar Error de Integridad", "cu20"),
    (1, "T02.3. Diseño Desbloquear Usuario", "t023"),
    (1, "T02.4. Diseño Modificar Usuario", "t024"),
    (1, "T02.5. Diseño Cambiar Clave", "t025"),
    (1, "T02.6. Diseño Logout", "t026"),
    (1, "T02.7. Diseño Activar/Desactivar Usuario", "t027"),
    (0, "T03. Encriptación", "t03"),
    (0, "T04. Perfiles", "t04"),
    (1, "T04.1. Crear Familia", "t041"),
    (1, "T04.2. Eliminar Familia", "t042"),
    (1, "T04.3. Crear Perfil", "t043"),
    (1, "T04.4. Eliminar Perfil", "t044"),
    (1, "T04.5. Agregar Permisos a Familia", "t045"),
    (1, "T04.6. Quitar Permisos a Familia", "t046"),
    (1, "T04.7. Agregar Permisos a Perfil", "t047"),
    (1, "T04.8. Quitar Permisos a Perfil", "t048"),
    (1, "T04.9. Agregar Familias a Perfil", "t049"),
    (1, "T04.10. Quitar Familias a Perfil", "t0410"),
    (0, "T05. Idiomas", "t05"),
    (0, "T06. Bitácora de Eventos", "t06"),
    (0, "T07. Gestión de Respaldos", "t07"),
    (0, "T08. Dígito Verificador", "t08"),
]

def main():
    fixed_path = r'c:\Users\usuario\OneDrive\Facultad\3° Año\2° Cuatrimestre\Trabajo de Diploma\Carpeta\Proyecto-Trabajo-de-Diploma\scratch\Proyecto_fixed.docx'
    orig_path = r'c:\Users\usuario\OneDrive\Facultad\3° Año\2° Cuatrimestre\Trabajo de Diploma\Carpeta\Proyecto-Trabajo-de-Diploma\scratch\Proyecto.docx'

    print("Loading documents...")
    doc = docx.Document(fixed_path)
    orig_doc = docx.Document(orig_path)

    print(f"Original paragraphs: {len(orig_doc.paragraphs)}")
    print(f"Fixed paragraphs:    {len(doc.paragraphs)}")
    print(f"Original tables:     {len(orig_doc.tables)}")
    print(f"Fixed tables:        {len(doc.tables)}")

    # 1. P24 check
    assert "Índice:" in doc.paragraphs[24].text, f"P24 is not Índice: {doc.paragraphs[24].text}"
    print("[OK] P24 is Índice:")

    # 2. P25 check
    assert doc.paragraphs[25].text.strip() == "", f"P25 is not empty: {doc.paragraphs[25].text}"
    print("[OK] P25 is empty separator")

    # 3. Check all 54 index entries
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    for idx, (expected_level, expected_text, expected_anchor) in enumerate(INDEX_ENTRIES):
        p_idx = 26 + idx
        p = doc.paragraphs[p_idx]
        assert p.text == expected_text, f"At P{p_idx}: expected '{expected_text}', got '{p.text}'"

        # Check hyperlink element
        hl = p._p.xpath('.//w:hyperlink')
        assert len(hl) == 1, f"At P{p_idx}: expected 1 hyperlink, found {len(hl)}"
        anchor = hl[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}anchor')
        assert anchor == expected_anchor, f"At P{p_idx}: expected anchor '{expected_anchor}', got '{anchor}'"

        # Check indentation
        ind = p._p.xpath('.//w:ind')
        if expected_level == 0:
            assert len(ind) == 0, f"At P{p_idx}: expected no ind, got {ind}"
        elif expected_level == 1:
            assert len(ind) == 1 and ind[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine') == '720', f"At P{p_idx}: expected firstLine=720"
        elif expected_level == 2:
            assert len(ind) == 1 and ind[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}firstLine') == '1440', f"At P{p_idx}: expected firstLine=1440"

    print("[OK] All 54 index entries verified with exact text, hyperlinks, anchors, and indentation!")

    # 4. Check page break paragraph at P80
    p80 = doc.paragraphs[26 + len(INDEX_ENTRIES)]
    br = p80._p.xpath('.//w:br')
    assert len(br) == 1 and br[0].attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}type') == 'page', f"P80 is not page break: {p80._p.xml}"
    print("[OK] Page break paragraph cleanly placed after index entries")

    # 5. Check G01 starts right after page break
    p81 = doc.paragraphs[26 + len(INDEX_ENTRIES) + 1]
    assert "G01. Propósito y Motivación" in p81.text, f"P81 is not G01: {p81.text}"
    bm = p81._p.xpath('.//w:bookmarkStart')
    assert any(b.attrib.get('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}name') == 'g01' for b in bm), "g01 bookmark missing on P81"
    print("[OK] G01 starts immediately after page break with bookmark g01")

    # 6. Check that the rest of the document matches original exactly
    # Original body starts at orig_doc.paragraphs[67]
    # Fixed body starts at doc.paragraphs[81]
    orig_body_paras = orig_doc.paragraphs[67:]
    fixed_body_paras = doc.paragraphs[81:]
    assert len(orig_body_paras) == len(fixed_body_paras), f"Body paragraph count mismatch: orig={len(orig_body_paras)}, fixed={len(fixed_body_paras)}"
    print(f"[OK] Body paragraph count matches exactly: {len(fixed_body_paras)} paragraphs")

    mismatches = 0
    for i, (orig_p, fix_p) in enumerate(zip(orig_body_paras, fixed_body_paras)):
        if orig_p.text != fix_p.text:
            print(f"Mismatch at body para {i}: orig='{orig_p.text[:30]}' != fix='{fix_p.text[:30]}'")
            mismatches += 1
            if mismatches > 5:
                break

    assert mismatches == 0, f"Found {mismatches} mismatches in document body!"
    print("[OK] 100% of body paragraphs match the original document text identically!")

    # 7. Check bookmark targets exist in document
    all_bms = set(doc._body._body.xpath('.//w:bookmarkStart/@w:name'))
    missing_bms = []
    for _, _, anchor in INDEX_ENTRIES:
        if anchor not in all_bms:
            missing_bms.append(anchor)
    assert len(missing_bms) == 0, f"Missing bookmarks: {missing_bms}"
    print("[OK] All 54 anchor bookmarks exist in the document body!")

    print("\nALL VALIDATIONS PASSED PERFECTLY!")

if __name__ == '__main__':
    main()
