import docx

doc = docx.Document(r"scratch/Proyecto.docx")
print("Total paragraphs:", len(doc.paragraphs))

with open("scratch/doc_structure.txt", "w", encoding="utf-8") as f:
    for i in range(65, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        txt = p.text.strip()
        if not txt:
            continue
        # Check if heading-like
        is_heading = False
        if p.style.name.startswith("Heading") or p.style.name in ["Title", "Subtitle"]:
            is_heading = True
        elif any(txt.startswith(pref) for pref in ["G0", "N0", "T0", "A0", "R0", "PN", "CU", "CUN", "RF"]):
            is_heading = True
        elif len(txt) < 80 and any(kw in txt for kw in ["Diagrama", "Especificación", "Especificacion", "Modelo", "GUI", "Requisitos", "Arquitectura", "Usuarios", "Perfiles", "Idioma", "Bitácora", "Bitacora", "Dígito", "Digito", "Respaldos", "Roles"]):
            is_heading = True
        
        if is_heading:
            f.write(f"P{i} [{p.style.name}]: {txt}\n")

print("Wrote doc_structure.txt")
