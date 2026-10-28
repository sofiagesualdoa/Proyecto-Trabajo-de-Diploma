import docx

doc = docx.Document(r"C:\Users\usuario\.gemini\antigravity\brain\6d6f5d5a-4cad-4d65-af9e-bdf70f4db580\.user_uploaded\media_1793025563617_a3f2539d.docx")
for i in range(1081, 1120):
    if i < len(doc.paragraphs):
        p = doc.paragraphs[i]
        if p.text.strip():
            print(f"[{i}] {p.text.strip()}")
