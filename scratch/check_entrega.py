import os
import docx

uploaded_dir = r"C:\Users\usuario\.gemini\antigravity\brain\6d6f5d5a-4cad-4d65-af9e-bdf70f4db580\.user_uploaded"
for f in os.listdir(uploaded_dir):
    if f.endswith(".docx"):
        full = os.path.join(uploaded_dir, f)
        print(f"=== {f} ===")
        doc = docx.Document(full)
        print("Total paragraphs:", len(doc.paragraphs))
        print("Total tables:", len(doc.tables))
        for i, p in enumerate(doc.paragraphs[:25]):
            if p.text.strip():
                print(f"  P{i}: {p.text.strip()[:100]}")
