import docx

doc = docx.Document(r"scratch/Proyecto.docx")
body = doc._body._body

p24 = doc.paragraphs[24]._p
p25 = doc.paragraphs[25]._p
p26 = doc.paragraphs[26]._p
p66 = doc.paragraphs[66]._p
p67 = doc.paragraphs[67]._p

idx24 = body.index(p24)
idx25 = body.index(p25)
idx26 = body.index(p26)
idx66 = body.index(p66)
idx67 = body.index(p67)

print(f"body indices: P24={idx24}, P25={idx25}, P26={idx26}, P66={idx66}, P67={idx67}")
print(f"Number of elements between P25 and P67: {idx67 - idx25 - 1}")
