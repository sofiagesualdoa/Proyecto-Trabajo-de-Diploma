import docx

doc = docx.Document(r"C:\Users\usuario\.gemini\antigravity\brain\6d6f5d5a-4cad-4d65-af9e-bdf70f4db580\.user_uploaded\media_1793025563617_a3f2539d.docx")
print("Total paragraphs:", len(doc.paragraphs))
for i in range(len(doc.paragraphs)-200, len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if t.startswith("N02") or t.startswith("CU") or "Instalad" in t or "Ayuda" in t or "Serializ" in t or "Cambios" in t or "Reporte" in t or "Reinstal" in t:
        print(f"[{i}] {t[:100]}")
