import docx

doc = docx.Document(r"C:\Users\usuario\.gemini\antigravity\brain\6d6f5d5a-4cad-4d65-af9e-bdf70f4db580\.user_uploaded\media_1793025563617_a3f2539d.docx")
for i in range(950, len(doc.paragraphs)):
    t = doc.paragraphs[i].text.strip()
    if t and len(t) > 3:
        if t.startswith("A0") or t.startswith("N0") or t.startswith("CU") or t.startswith("Reporte") or t.startswith("G") or t.startswith("Manual") or t.startswith("Instal"):
            print(f"[{i}] {t}")
