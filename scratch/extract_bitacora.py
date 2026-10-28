import docx

doc = docx.Document(r'C:\Users\usuario\.gemini\antigravity\brain\6d6f5d5a-4cad-4d65-af9e-bdf70f4db580\.user_uploaded\media_1793025563617_a3f2539d.docx')
with open(r'c:\Users\usuario\OneDrive\Facultad\3° Año\2° Cuatrimestre\Trabajo de Diploma\Carpeta\Proyecto-Trabajo-de-Diploma\scratch\classmate_bitacora.txt', 'w', encoding='utf-8') as out:
    for i in range(972, min(1035, len(doc.paragraphs))):
        out.write(f"P{i}: {doc.paragraphs[i].text}\n")
print("Saved classmate_bitacora.txt")
